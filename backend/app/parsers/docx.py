"""DOCX 简历解析：把 .docx 字节流抽取为纯文本。

输入约定为 **Word 文字版**（需求 D1）——不支持扫描件与 OCR，
也不支持 .doc（OLE2）与 PDF，二者在接口层就被拦下。

解析的目标不是"还原版式"，而是**给模型一段干净、按原文顺序排列的文字**。
因此表格被拍平成行、页眉页脚合并到文末，位置信息会丢失——这些都会通过
``warnings`` 告诉用户，由前端回显确认后再提交分析（见前置文档 2.3 与第十节）。
"""

from __future__ import annotations

import io
from typing import Iterator

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


class DocxParseError(Exception):
    """文档无法解析（损坏、加密、非 docx 等）。接口层据此返回 400。"""


def extract_text(data: bytes) -> tuple[str, list[str]]:
    """解析 .docx，返回 ``(纯文本, 告警列表)``。

    告警用于提示"解析到了但可能不完整"的地方，前端应把它们显示给用户，
    提醒其确认内容无误——Word 排版复杂时解析结果难免有出入。
    """
    document = _open(data)
    warnings: list[str] = []

    lines: list[str] = []
    table_count = 0
    for block in _iter_block_items(document):
        if isinstance(block, Table):
            table_count += 1
            lines.extend(_table_to_lines(block))
        else:
            text = _clean(block.text)
            if text:
                lines.append(text)

    if table_count:
        warnings.append(f"文档中包含 {table_count} 个表格，已按行合并为文本")

    header_footer = _header_footer_lines(document)
    if header_footer:
        lines.extend(header_footer)
        warnings.append("已合并页眉/页脚中的文本")

    textbox = _dedup_textbox_lines(_textbox_lines(document), set(lines))
    if textbox:
        lines.extend(textbox)
        warnings.append("检测到文本框，其内容已追加到文末，位置可能与原文不一致")

    text = "\n".join(lines)

    if not text:
        warnings.append("未从文档中解析出任何文本，请确认上传的是文字版简历")
    elif len(text) < 50:
        warnings.append(f"解析出的文本较少（{len(text)} 字），请确认内容完整")

    return text, warnings


# ── 内部实现 ────────────────────────────────────


def _clean(text: str) -> str:
    """把一段文本里的换行与连续空白压成单个空格。"""
    return " ".join(text.split())


def _open(data: bytes) -> DocxDocument:
    if not data:
        raise DocxParseError("文件内容为空")
    try:
        return Document(io.BytesIO(data))
    except Exception as exc:  # python-docx 对损坏文件抛的异常类型不统一
        raise DocxParseError("文档无法解析，请确认是有效的 .docx 文字版简历") from exc


def _iter_block_items(document: DocxDocument) -> Iterator[Paragraph | Table]:
    """按文档顺序产出段落与表格。

    python-docx 的 ``.paragraphs`` 与 ``.tables`` 是两个独立的列表，
    直接用会丢失二者在原文中的先后顺序——简历里"先写一段说明再放表格"很常见，
    顺序错了模型读到的上下文就是乱的。
    """
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _table_to_lines(table: Table) -> list[str]:
    """把表格拍平成「单元格 | 单元格」的行，便于模型按行理解。"""
    lines: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            text = _clean(cell.text)
            # 横向合并的单元格会在每一列上重复出现，跳过连续重复项
            if text and (not cells or cells[-1] != text):
                cells.append(text)
        if cells:
            lines.append(" | ".join(cells))
    return lines


def _header_footer_lines(document: DocxDocument) -> list[str]:
    """取页眉页脚文本。联系方式常被放在页眉里，漏掉会丢失电话与邮箱。"""
    lines: list[str] = []
    for section in document.sections:
        for part in (section.header, section.footer):
            if part is None:
                continue
            for paragraph in part.paragraphs:
                text = _clean(paragraph.text)
                if text:
                    lines.append(text)
    return lines


def _textbox_lines(document: DocxDocument) -> list[str]:
    """取文本框内的文字。

    python-docx 不暴露文本框，只能直接翻 XML 找 ``w:txbxContent``。
    """
    lines: list[str] = []
    for textbox in document.element.body.iter(qn("w:txbxContent")):
        for paragraph in textbox.iter(qn("w:p")):
            text = _clean("".join(node.text or "" for node in paragraph.iter(qn("w:t"))))
            if text:
                lines.append(text)
    return lines


def _dedup_textbox_lines(textbox_lines: list[str], body_lines: set[str]) -> list[str]:
    """去掉文本框内部的重复行，以及与正文重复的行。

    mc:AlternateContent 会把同一个文本框在 Choice 与 Fallback 里各存一份，
    不去重会把内容读成两遍。
    """
    result: list[str] = []
    seen: set[str] = set()
    for line in textbox_lines:
        if line in body_lines or line in seen:
            continue
        seen.add(line)
        result.append(line)
    return result
